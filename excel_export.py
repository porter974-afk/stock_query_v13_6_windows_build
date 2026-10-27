"""Generate a desktop-compatible report entirely in memory."""
from io import BytesIO
import re

import pandas as pd
from openpyxl.styles import Alignment, Font

from stock_core import taipei_today


def format_price(value):
    return f"{float(value):.2f}".rstrip("0").rstrip(".") if pd.notna(value) else ""


def display_frame(data, legacy_first=False):
    frame = data.copy()
    weekdays = ["週一", "週二", "週三", "週四", "週五", "週六", "週日"]
    frame["日期"] = frame["日期"].map(lambda d: f"{d:%Y/%m/%d} ({weekdays[d.weekday()]})")
    for col in ["開盤價", "最高價", "最低價", "收盤價"]:
        frame[col] = frame[col].map(format_price)
    frame["漲跌價"] = data["漲跌價"].map(lambda v: ("0" if legacy_first else "—") if pd.isna(v) else ("0" if v == 0 else ("▲" if v > 0 else "▼") + format_price(abs(v))))
    frame["漲跌幅(%)"] = data["漲跌幅(%)"].map(lambda v: ("0.00%" if legacy_first else "—") if pd.isna(v) else f"{abs(v):.2f}%")
    return frame[["日期", "成交量", "開盤價", "最高價", "最低價", "收盤價", "漲跌價", "漲跌幅(%)"]]


def export_excel(result, range_label):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        display_frame(result.data, legacy_first=True).to_excel(writer, sheet_name="報表", index=False, startrow=1)
        ws = writer.sheets["報表"]
        ws.freeze_panes = "A3"
        ws.cell(1, 1, f"標的：{result.name} ({result.symbol}) | 查詢期間：{range_label} | 顏色區分漲跌 (無正負號)")
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        for row, diff in enumerate(result.data["漲跌價"], start=3):
            ws.cell(row, 2).number_format = "#,##0"
            if pd.notna(diff) and diff != 0:
                for col in (7, 8):
                    ws.cell(row, col).font = Font(color="FF0000" if diff > 0 else "008000", bold=True)
        for col, width in {"A": 24, "B": 15, "C": 14, "D": 14, "E": 14, "F": 14, "G": 12, "H": 14}.items():
            ws.column_dimensions[col].width = width
    safe_name = re.sub(r'[\\/*?:"<>|\x00-\x1f]', "", result.name)
    return buffer.getvalue(), f"{safe_name}_{range_label}_{taipei_today():%Y%m%d}.xlsx"
