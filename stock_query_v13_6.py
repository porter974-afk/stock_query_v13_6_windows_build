# -*- coding: utf-8 -*-
import sys
import os
import ctypes
import io
import re
from datetime import date, datetime, timedelta
import tkinter as tk
from tkinter import messagebox, ttk

# 1. 【閃電響應】
print("\n" + "="*45)
print("🚀 [V13.6 版] 台美股價查詢系統啟動中...")
print("="*45 + "\n")

# 修正輸出重導向
if sys.stdout is None: sys.stdout = io.StringIO()
if sys.stderr is None: sys.stderr = io.StringIO()

def get_resource_path(relative_path):
    try:
        base_path = sys._MEIPASS 
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def show_splash():
    splash = tk.Tk()
    splash.title("Loading")
    w, h = 350, 130
    sw, sh = splash.winfo_screenwidth(), splash.winfo_screenheight()
    splash.geometry(f"{w}x{h}+{int((sw-w)/2)}+{int((sh-h)/2)}")
    splash.overrideredirect(True)
    splash.attributes("-topmost", True)
    frame = tk.Frame(splash, relief='ridge', borderwidth=4, bg="#f8f9fa")
    frame.pack(fill='both', expand=True)
    label = tk.Label(frame, text="\n🌍 數據精鍊美化中...\n\n請稍候，報表工具即將開啟", 
                     font=("微軟正黑體", 10, "bold"), bg="#f8f9fa", fg="#333")
    label.pack(expand=True)
    splash.update()
    return splash

splash_win = show_splash()

try:
    import yfinance as yf
    import pandas as pd
    from openpyxl.styles import Font, Alignment
except Exception as e:
    print(f"致命錯誤：套件載入失敗 - {e}")
    sys.exit(1)

def hide_console():
    try:
        whnd = ctypes.windll.kernel32.GetConsoleWindow()
        if whnd != 0:
            ctypes.windll.user32.ShowWindow(whnd, 0)
    except: pass

hide_console()
splash_win.destroy()

# ---------- 核心功能 ----------

def smart_format_price(val):
    """ 智能顯示價格：整數去小數，小數留兩位 """
    try:
        f_val = float(val)
        if f_val == 0: return "0"
        if f_val.is_integer(): return str(int(f_val))
        return "{:.2f}".format(f_val).rstrip('0').rstrip('.')
    except:
        return val

def get_stock_id_smart(search_text):
    search_text = search_text.strip().upper()
    if not search_text: return None
    if re.match(r'^[A-Z]+$', search_text): return search_text
    if search_text.isdigit(): return search_text
    hot_map = {"日月光": "3711", "鴻海": "2317", "台積電": "2330", "輝達": "NVDA", "特斯拉": "TSLA"}
    if search_text in hot_map: return hot_map[search_text]
    try:
        import twstock
        for code, info in twstock.codes.items():
            if search_text == info.name or search_text in info.name: return code
    except: pass
    return search_text

def process_single_stock(user_input, output_dir, start_date=None, end_date=None, range_label="近一個月"):
    symbol = get_stock_id_smart(user_input)
    if not symbol: return False, "無效代碼"
    
    df = pd.DataFrame()
    display_name = symbol 
    
    ticker_list = [f"{symbol}.TW", f"{symbol}.TWO"] if symbol.isdigit() else [symbol]
        
    for ticker in ticker_list:
        try:
            stock_obj = yf.Ticker(ticker)
            if start_date is None:
                df = stock_obj.history(period="1mo", auto_adjust=False)
            else:
                # yfinance 的 end 為不包含當天，故增加一天
                df = stock_obj.history(start=start_date.isoformat(),
                                       end=(end_date + timedelta(days=1)).isoformat(),
                                       auto_adjust=False)
            if not df.empty:
                try:
                    import twstock
                    if symbol in twstock.codes: display_name = twstock.codes[symbol].name
                except: pass
                break
        except: continue
        
    if df.empty: return False, f"{symbol} 無數據"
    
    # --- 數據處理 ---
    df = df.reset_index()
    if 'Date' in df.columns:
        df['Date'] = pd.to_datetime(df['Date']).dt.tz_localize(None)
        weekday_map = {0:"週一", 1:"週二", 2:"週三", 3:"週四", 4:"週五", 5: "週六", 6: "週日"}
        df['日期'] = df['Date'].apply(lambda d: f"{d.strftime('%Y/%m/%d')} ({weekday_map[d.weekday()]})")
        df = df.drop(columns=['Date'])
    
    df = df.rename(columns={"Open":"開盤價", "High":"最高價", "Low":"最低價", "Close":"收盤價", "Volume":"成交量"})
    
    # 漲跌核心計算
    df["RawDiff"] = df["收盤價"].diff()
    df["RawPct"] = (df["收盤價"].pct_change() * 100).round(2)
    
    # 格式化開高低收
    for col in ["開盤價", "最高價", "最低價", "收盤價"]:
        df[col] = df[col].apply(smart_format_price)
    
    # 格式化漲跌價 (▲/▼ + 無負號)
    def format_diff_visual(val):
        if pd.isna(val) or val == 0: return "0"
        icon = "▲" if val > 0 else "▼"
        return f"{icon}{smart_format_price(abs(val))}"
    
    # 格式化漲跌幅 (無符號 + 兩位小數 + %)
    def format_pct_visual(val):
        if pd.isna(val) or val == 0: return "0.00%"
        # 絕對值化，移除 +/- 符號
        return "{:.2f}%".format(abs(val))
        
    df["漲跌價"] = df["RawDiff"].apply(format_diff_visual)
    df["漲跌幅(%)"] = df["RawPct"].apply(format_pct_visual)
    
    # 欄位排序
    df = df[['日期', '成交量', '開盤價', '最高價', '最低價', '收盤價', '漲跌價', '漲跌幅(%)', 'RawDiff']]
    
    # --- 儲存與對齊 ---
    safe_name = re.sub(r'[\\/*?:"<>|]', "", display_name)
    file_path = os.path.join(output_dir, f"{safe_name}_{range_label}_{date.today().strftime('%Y%m%d')}.xlsx")
    
    writer = pd.ExcelWriter(file_path, engine="openpyxl")
    df.drop(columns=['RawDiff']).to_excel(writer, sheet_name="報表", index=False, startrow=1)
    ws = writer.sheets["報表"]
    ws.freeze_panes = "A3"
    
    left_align = Alignment(horizontal='left', vertical='center')
    red_f = Font(color="FF0000", bold=True)
    grn_f = Font(color="008000", bold=True)
    
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = left_align
            
    for i, r_idx in enumerate(range(3, ws.max_row + 1)):
        ws.cell(r_idx, 2).number_format = '#,##0' # 成交量
        
        diff_val = df.iloc[i]["RawDiff"]
        if pd.notna(diff_val):
            # 漲跌價與漲跌幅著色
            if diff_val > 0:
                ws.cell(r_idx, 7).font = red_f
                ws.cell(r_idx, 8).font = red_f
            elif diff_val < 0:
                ws.cell(r_idx, 7).font = grn_f
                ws.cell(r_idx, 8).font = grn_f

    ws.cell(1, 1).value = f"標的：{display_name} ({symbol}) | 查詢期間：{range_label} | 顏色區分漲跌 (無正負號)"
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12
    writer.close()
    return True, display_name

def get_selected_range():
    mode = mode_var.get()
    today = date.today()
    if mode == "最近一個月":
        return None, None, "近一個月"
    if mode == "指定年份":
        year = int(year_var.get())
        if year > today.year:
            raise ValueError("年份不能晚於今年")
        return date(year, 1, 1), min(date(year, 12, 31), today), str(year)
    try:
        start = datetime.strptime(start_var.get().strip(), "%Y-%m-%d").date()
        end = datetime.strptime(end_var.get().strip(), "%Y-%m-%d").date()
    except ValueError:
        raise ValueError("請輸入正確日期，格式為 YYYY-MM-DD，例如 2026-01-01")
    if start > end:
        raise ValueError("開始日期不能晚於結束日期")
    if start > today or end > today:
        raise ValueError("查詢日期不能晚於今天")
    return start, end, f"{start:%Y%m%d}-{end:%Y%m%d}"


def update_date_controls(*_):
    mode = mode_var.get()
    year_box.configure(state="readonly" if mode == "指定年份" else "disabled")
    state = "normal" if mode == "自訂日期區間" else "disabled"
    start_entry.configure(state=state)
    end_entry.configure(state=state)


def batch_run():
    raw_input = entry.get().strip()
    if not raw_input:
        messagebox.showwarning("提醒", "請先輸入股票代碼")
        return
    try:
        start_date, end_date, range_label = get_selected_range()
    except ValueError as exc:
        messagebox.showerror("日期錯誤", str(exc))
        return
    items = re.split(r'[，,]+', raw_input)
    base = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
    out = os.path.join(base, "output")
    os.makedirs(out, exist_ok=True)
    
    success = 0
    failures = []
    for i in items:
        t = i.strip()
        if not t: continue
        status_var.set(f"處理中: {t}"); root.update()
        try:
            res, info = process_single_stock(t, out, start_date, end_date, range_label)
        except Exception as exc:
            res, info = False, str(exc)
        if res:
            success += 1
        else:
            failures.append(f"{t}：{info}")
    
    status_var.set("完成")
    if success:
        try:
            if sys.platform == "win32":
                os.startfile(out)
            elif sys.platform == "darwin":
                import subprocess
                subprocess.Popen(["open", out])
            else:
                import subprocess
                subprocess.Popen(["xdg-open", out])
        except Exception:
            pass
    detail = "\n".join(failures[:8])
    messagebox.showinfo("完成", f"成功處理 {success} 筆股票。" + (f"\n失敗 {len(failures)} 筆：\n{detail}" if failures else ""))

# --- UI ---
root = tk.Tk()
root.title("不專業股價報表 V13.6版")
root.geometry("470x390")
try: root.iconbitmap(get_resource_path("stock_1.ico"))
except: pass

tk.Label(root, text="輸入股票代碼 (台美股皆可)", font=("微軟正黑體", 11, "bold"), pady=15).pack()
entry = tk.Entry(root, width=35, font=("Arial", 11), justify='center'); entry.pack()
entry.insert(0, "2330, NVDA, 00905")
# 日期查詢選單
range_frame = ttk.LabelFrame(root, text="查詢日期範圍", padding=12)
range_frame.pack(fill="x", padx=18, pady=(15, 0))
mode_var = tk.StringVar(value="最近一個月")
mode_box = ttk.Combobox(range_frame, textvariable=mode_var,
                        values=["最近一個月", "指定年份", "自訂日期區間"], state="readonly", width=18)
mode_box.grid(row=0, column=1, sticky="w", pady=4)
ttk.Label(range_frame, text="查詢模式：").grid(row=0, column=0, sticky="e")
year_var = tk.StringVar(value=str(date.today().year))
year_box = ttk.Combobox(range_frame, textvariable=year_var,
                        values=[str(y) for y in range(date.today().year, 1989, -1)],
                        state="disabled", width=16)
ttk.Label(range_frame, text="選擇年份：").grid(row=1, column=0, sticky="e")
year_box.grid(row=1, column=1, sticky="w", pady=4)
start_var = tk.StringVar(value=f"{date.today().year}-01-01")
end_var = tk.StringVar(value=date.today().isoformat())
ttk.Label(range_frame, text="開始日期：").grid(row=2, column=0, sticky="e")
start_entry = ttk.Entry(range_frame, textvariable=start_var, width=19, state="disabled")
start_entry.grid(row=2, column=1, sticky="w", pady=4)
ttk.Label(range_frame, text="結束日期：").grid(row=3, column=0, sticky="e")
end_entry = ttk.Entry(range_frame, textvariable=end_var, width=19, state="disabled")
end_entry.grid(row=3, column=1, sticky="w", pady=4)
ttk.Label(range_frame, text="日期格式：YYYY-MM-DD", foreground="gray").grid(row=4, column=1, sticky="w")
mode_box.bind("<<ComboboxSelected>>", update_date_controls)

tk.Button(root, text="一鍵生成報表", command=batch_run, bg="#28a745", fg="white", font=("微軟正黑體", 10, "bold"), width=20, height=2).pack(pady=20)
status_var = tk.StringVar(value="準備就緒")
tk.Label(root, textvariable=status_var, fg="gray", font=("微軟正黑體", 9)).pack(side=tk.BOTTOM, pady=5)
root.mainloop()
