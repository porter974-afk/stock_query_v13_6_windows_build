from datetime import date
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from excel_export import export_excel
from stock_core import StockResult, fetch_stock, parse_symbols, prepare_history, resolve_symbol, selected_range


def history():
    return pd.DataFrame({"Open": [100, 110, 99], "High": [101, 111, 100], "Low": [99, 109, 98], "Close": [100, 110, 99], "Volume": [1000, 2000, 3000]}, index=pd.date_range("2026-01-05", periods=3, tz="Asia/Taipei"))


class CoreTests(unittest.TestCase):
    def test_symbols_and_dates(self):
        self.assertEqual(parse_symbols("00905，nvda 00905"), ["00905", "NVDA"])
        self.assertEqual(resolve_symbol("台積電"), "2330")
        self.assertEqual(resolve_symbol("00905"), "00905")
        today = date(2026, 10, 28)
        self.assertEqual(selected_range("指定年份", 2026, today=today)[1], today)
        self.assertEqual(selected_range("指定年份", 2024, today=today)[1], date(2024, 12, 31))
        for start, end in [(today, date(2026, 1, 1)), (today, date(2026, 10, 29))]:
            with self.assertRaises(ValueError):
                selected_range("自訂日期區間", start=start, end=end, today=today)

    def test_numeric_and_excel(self):
        frame = prepare_history(history())
        self.assertEqual(frame["漲跌價"].iloc[1:].tolist(), [10, -11])
        self.assertEqual(frame["漲跌幅(%)"].iloc[1:].tolist(), [10, -10])
        self.assertIsNone(frame["日期"].dt.tz)
        content, filename = export_excel(StockResult("2330", "2330.TW", "台積電", frame), "2026")
        ws = load_workbook(BytesIO(content)).active
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertEqual(ws.freeze_panes, "A3")
        self.assertEqual(ws["B3"].value, 1000)
        self.assertEqual(ws["G4"].value, "▲10")
        self.assertEqual(ws["G5"].value, "▼11")
        self.assertEqual(ws["G4"].font.color.rgb, "00FF0000")
        self.assertEqual(ws["G5"].font.color.rgb, "00008000")

    @patch("stock_core.yf.Ticker")
    def test_fallback_and_inclusive_end(self, ticker):
        ticker.return_value.history.side_effect = [pd.DataFrame(), history()]
        result = fetch_stock("6488", date(2026, 1, 5), date(2026, 1, 7))
        self.assertEqual(result.ticker, "6488.TWO")
        self.assertEqual(ticker.return_value.history.call_args.kwargs["end"], "2026-01-08")
        self.assertFalse(ticker.return_value.history.call_args.kwargs["auto_adjust"])

    @patch("stock_core.yf.Ticker")
    def test_all_candidates_empty(self, ticker):
        ticker.return_value.history.return_value = pd.DataFrame()
        with self.assertRaisesRegex(ValueError, "無法取得股價"):
            fetch_stock("00000")

    def test_page_query_and_persistence(self):
        result = StockResult("2330", "2330.TW", "台積電", prepare_history(history()))
        with patch("stock_core.fetch_stock", return_value=result):
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "streamlit_app.py")).run()
            self.assertFalse(app.exception)
            app.text_input[0].set_value("2330")
            app.button[0].click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.session_state["query"]["results"]), 1)

    def test_page_partial_failure_and_validation(self):
        result = StockResult("NVDA", "NVDA", "NVDA", prepare_history(history()))
        with patch("stock_core.fetch_stock", side_effect=[result, ValueError("無效代碼")]):
            app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "streamlit_app.py")).run()
            app.text_input[0].set_value("NVDA, INVALID_TEST")
            app.button[0].click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.session_state["query"]["results"]), 1)
            self.assertEqual(len(app.session_state["query"]["failures"]), 1)
            app.text_input[0].set_value("")
            app.button[0].click().run()
            self.assertTrue(app.error)
            self.assertFalse(app.exception)
            app.multiselect[0].set_value(["最高價", "最低價"]).run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.session_state["query"]["results"]), 1)


if __name__ == "__main__":
    unittest.main()
